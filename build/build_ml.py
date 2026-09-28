"""V5 builder - part 2: ML_Tank_Training (deduped, fully formula-traceable feature store)."""
import pandas as pd, numpy as np, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL

FONT='Arial'
# P1-5: row bounds are DERIVED from the source, never typed. build_core reads archive.csv
# and exposes N, so adding source rows needs no edit here. HANDOFF gotcha 6.
# Clean_Data column letters (must match build_core.HEADERS order)
from build_core import C as CD, HEADERS as CDH, N as N_CLEAN
CF,CLST=2,N_CLEAN+1                        # Clean_Data data-row span

# ---------------------------------------------------------------- ML layout
# (header, group, kind) kind: pull|calc|target|ref
SPEC=[]
def add(h,group,kind,f=None,fmt=None,note=''):
    SPEC.append(dict(h=h,group=group,kind=kind,f=f,fmt=fmt,note=note))

# identity
add('Tank Key','Identity','key',None,None,'Unique row id: Quote-Rev-TankName (dedup-suffixed)')
add('Src Row','Identity','calc',lambda r,M:f'=MATCH($A{r},Clean_Data!${CD["Tank Key"]}${CF}:${CD["Tank Key"]}${CLST},0)+1',
    '0','Row of this tank in Clean_Data - every pull below indexes this row')
PULL=[('QuoteNumber','QuoteNumber'),('RevisionNumber','RevisionNumber'),('Quote-Rev','Quote-Rev'),
      ('Tank Name','Tank Name'),('Bid Type','Bid Type'),('Status','Status'),('Outcome','Outcome'),
      ('Revision Count','Revision Count'),('Due Date','Due Date'),('Due Year','Due Year'),
      ('Country','Country'),('State','State'),('City','City'),
      # P0-2: normalized geography. Country/State above are the raw source values and are
      # dirty (208 State values mixing US codes, Mexican codes, cities and countries; the
      # Country column itself is wrong on 37 rows). Use these three for any geographic
      # grouping. See geo_crosswalk.py.
      ('Country Normalized','Country Normalized'),('State Normalized','State Normalized'),
      ('Region Code','Region Code'),
      ('Sales Manager','Sales Manager'),
      ('Company Name','Company Name'),('Customer Name','Customer Name'),
      ('Quote Group ID','Quote Group ID'),('Quote Variant Type','Quote Variant Type'),
      ('Is Scope Variant','Is Scope Variant'),('Effective Revision','Effective Revision'),
      ('Job Tank Key','Job Tank Key')]
for h,src in PULL:
    # Due Date pulls a DATEVALUE serial; without an explicit date format it renders as
    # "45596" and pandas reads it as int64. The delivered V5 only looked right because
    # LibreOffice silently inferred the format during recalc - Excel does not.
    add(h,'Identity','pull',src,'yyyy-mm-dd' if h=='Due Date' else None)

# --- raw model inputs (what the web app supplies)
INPUTS=[('Material','Material'),('Use Type','Use Type'),('Deck Style','Deck Style'),
        ('Floor Style','Floor Style'),('Wage Type','Wage Type'),('Diameter (ft)','Diameter (ft)'),
        ('Height (ft)','Height (ft)'),('Freeboard (in)','Freeboard (in)'),
        ('Capacity Unit','Capacity Unit'),('Capacity Value','Capacity Value'),('Quantity','Quantity'),
        # P1-6: parsed from the tank NAME, not from any price - safe as model features,
        # unlike scope_class / is_insulated which are derived from the targets.
        ('Name Has Option','Name Has Option'),('Name Partial Scope','Name Partial Scope'),
        ('Ss','Ss'),('S1','S1'),('Miles TBT','Miles TBT'),('Miles GT','Miles GT'),
        ('Margin (%)','Margin (%)'),('Contingency (%)','Contingency (%)')]
for h,src in INPUTS:
    add(h,'Input (from web app)','pull',src)

# --- targets
TARGETS=[('Y Material','Material Price'),('Y Fabrication','Fabrication Price'),
         ('Y Construction','Construction Price'),('Y Insul Material','Insul Material Price'),
         ('Y Insul Construction','Insul Construction Price'),('Y Bucket Sum','Bucket Sum')]
for h,src in TARGETS:
    add(h,'TARGET','target',src,'$#,##0')
# reference-only money (computed by the estimator, never predicted)
for h,src in [('Ref Freight','Freight Price'),('Ref Tax','Total Tax'),('Ref Proposal Total','Proposal Total'),('Ref Total Price','Total Price')]:
    add(h,'Reference (not a target)','pull',src,'$#,##0')

L={}                                        # filled after layout is fixed
def col(h): return L[h]
ML_FIRST=3; ML_LAST=[0]                     # ML_LAST[0] set at build time
def mrng(h): return f"${col(h)}${ML_FIRST}:${col(h)}${ML_LAST[0]}"

# --- derived features (formulas over local columns)
D=[]
def dv(h,group,f,fmt=None,note=''):
    D.append((h,group,f,fmt,note)); add(h,group,'calc',None,fmt,note)

dv('geom_volume_cuft','Geometry',lambda r:f'=PI()*(${col("Diameter (ft)")}{r}/2)^2*${col("Height (ft)")}{r}','#,##0',
   'Cylinder volume from diameter & height')
dv('geom_volume_gal','Geometry',lambda r:f'=${col("geom_volume_cuft")}{r}*7.48052','#,##0','cu ft -> US gallons')
dv('shell_area_sqft','Geometry',lambda r:f'=PI()*${col("Diameter (ft)")}{r}*${col("Height (ft)")}{r}','#,##0',
   'Lateral shell area - primary steel/plate driver')
dv('floor_area_sqft','Geometry',lambda r:f'=PI()*(${col("Diameter (ft)")}{r}/2)^2','#,##0','Floor plate area')
dv('roof_slope_ratio','Geometry',lambda r:(
   f'=IFERROR(IFERROR(VALUE(MID(${col("Deck Style")}{r},FIND(":12",${col("Deck Style")}{r})-2,2)),'
   f'VALUE(MID(${col("Deck Style")}{r},FIND(":12",${col("Deck Style")}{r})-1,1)))/12,0)'),'0.000',
   'Parsed from deck style "n:12"; 0 when flat/open/dome. Reads TWO characters before the '
   'colon and falls back to one, so a two-digit pitch parses correctly. The earlier '
   'single-character form read "12:12" as 2:12 and "10:12" as 0 - no such deck style exists '
   'in the current data, so this changes no value today, but new data would have been wrong.')
dv('roof_area_sloped_sqft','Geometry',lambda r:f'=${col("floor_area_sqft")}{r}*SQRT(1+${col("roof_slope_ratio")}{r}^2)','#,##0',
   'Floor area adjusted for roof pitch')
dv('total_area_sqft','Geometry',lambda r:f'=${col("shell_area_sqft")}{r}+${col("floor_area_sqft")}{r}+${col("roof_area_sloped_sqft")}{r}','#,##0',
   'Total fabricated surface - strongest single price driver')
dv('circumference_ft','Geometry',lambda r:f'=PI()*${col("Diameter (ft)")}{r}','#,##0.0','Drives ring/seam welding length')
dv('aspect_ratio_hd','Geometry',lambda r:f'=IF(${col("Diameter (ft)")}{r}>0,${col("Height (ft)")}{r}/${col("Diameter (ft)")}{r},"")','0.000',
   'Slenderness H/D - tall narrow tanks cost more per gallon')
dv('diameter_sq','Geometry',lambda r:f'=${col("Diameter (ft)")}{r}^2','#,##0','Non-linear diameter term')
dv('height_sq','Geometry',lambda r:f'=${col("Height (ft)")}{r}^2','#,##0','Non-linear height term')
dv('d_x_h','Geometry',lambda r:f'=${col("Diameter (ft)")}{r}*${col("Height (ft)")}{r}','#,##0','Interaction term')
dv('freeboard_ft','Geometry',lambda r:f'=N(${col("Freeboard (in)")}{r})/12','0.00','Freeboard in feet (blank treated as 0)')
dv('freeboard_ratio','Geometry',lambda r:f'=IF(N(${col("Height (ft)")}{r})>0,${col("freeboard_ft")}{r}/${col("Height (ft)")}{r},"")','0.000',
   'Share of shell height that is freeboard')
dv('capacity_gal','Capacity',lambda r:f'=IF(AND(${col("Capacity Unit")}{r}="gal",'f'N(${col("Capacity Value")}{r})>0),${col("Capacity Value")}{r},"")','#,##0',
   'Usable capacity, gallons only')
dv('capacity_tons','Capacity',lambda r:f'=IF(AND(${col("Capacity Unit")}{r}="tons",'f'N(${col("Capacity Value")}{r})>0),${col("Capacity Value")}{r},"")','#,##0',
   'Usable capacity, tons only (mass - not convertible to gal)')
dv('is_ton_capacity','Capacity',lambda r:f'=IF(${col("Capacity Unit")}{r}="tons",1,0)','0','Silo/dry-bulk indicator')
dv('fill_ratio','Capacity',lambda r:f'=IF(AND(${col("geom_volume_gal")}{r}>0,${col("capacity_gal")}{r}<>""),${col("capacity_gal")}{r}/${col("geom_volume_gal")}{r},"")','0.000',
   'Usable/geometric volume - sanity check, typically ~0.93')
dv('material_grade_rank','Material',lambda r:f'=IF(${col("Material")}{r}="316SS",3,IF(${col("Material")}{r}="304SS",2,IF(${col("Material")}{r}="CS",1,0)))','0',
   'Ordinal grade CS<304SS<316SS')
dv('is_stainless','Material',lambda r:f'=IF(${col("material_grade_rank")}{r}>=2,1,0)','0','Stainless flag')
dv('material_cost_factor','Material',lambda r:f'=IF(${col("Material")}{r}="316SS",1.25,IF(${col("Material")}{r}="304SS",1.15,1))','0.00',
   'Relative $/lb factor vs carbon steel (assumption - see ML_Feature_Catalog)')
dv('material_area_intensity','Material',lambda r:f'=${col("shell_area_sqft")}{r}*${col("material_cost_factor")}{r}','#,##0',
   'Grade-weighted shell area')
dv('seismic_band','Seismic',lambda r:(f'=IF(${col("Ss")}{r}="","Unknown",IF(${col("Ss")}{r}<0.25,"1. Very Low",'
   f'IF(${col("Ss")}{r}<0.5,"2. Low",IF(${col("Ss")}{r}<1,"3. Moderate",IF(${col("Ss")}{r}<1.5,"4. High","5. Very High")))))'),None,
   'ASCE-style Ss banding')
dv('high_seismic','Seismic',lambda r:f'=IF(N(${col("Ss")}{r})>=0.5,1,0)','0','Ss>=0.5 flag')
dv('s1_ss_ratio','Seismic',lambda r:f'=IF(N(${col("Ss")}{r})>0,${col("S1")}{r}/${col("Ss")}{r},"")','0.000','Spectral shape')
dv('seismic_x_height','Seismic',lambda r:f'=N(${col("Ss")}{r})*${col("Height (ft)")}{r}','#,##0.0','Overturning/anchorage proxy')
dv('seismic_x_shell','Seismic',lambda r:f'=N(${col("Ss")}{r})*${col("shell_area_sqft")}{r}','#,##0','Seismic mass proxy')
dv('min_miles','Logistics',lambda r:(
   f'=IF(AND(N(${col("Miles TBT")}{r})>0,N(${col("Miles GT")}{r})>0),'
   f'MIN(${col("Miles TBT")}{r},${col("Miles GT")}{r}),'
   f'IF(N(${col("Miles TBT")}{r})>0,${col("Miles TBT")}{r},'
   f'IF(N(${col("Miles GT")}{r})>0,${col("Miles GT")}{r},"")))'),'#,##0',
   'Distance from the nearest plant, over the POSITIVE distances only. A zero plant distance '
   'means the figure was never computed, not that the site is at the plant: Miles GT is a '
   'literal 0 on 1,853 source rows and rising (1.4% of 2024 rows, 51.2% of 2026), while only '
   '26 rows have both distances zero. Taking a plain MIN let one zero drag the result to zero '
   'on 1,827 rows whose real distance ran to a median of 800 miles.')
dv('is_domestic','Logistics',lambda r:f'=IF(${col("Country Normalized")}{r}="US",1,0)','0',
   'US site. Keys off Country Normalized, not the raw Country column, which is wrong on 37 rows '
   '(e.g. US|Jalisco is Zapopan, Mexico; MO|MA is Uxbridge, Massachusetts).')
dv('is_international','Logistics',lambda r:f'=1-${col("is_domestic")}{r}','0','Non-US site')
dv('is_local_site','Logistics',lambda r:f'=IF(AND(${col("min_miles")}{r}<>"",${col("min_miles")}{r}<50),1,0)','0',
   'Within 50 mi of a plant. Zero when the distance is unknown rather than assuming local.')
dv('distance_band','Logistics',lambda r:(f'=IF(${col("min_miles")}{r}="","Unknown",'
   f'IF(${col("min_miles")}{r}<100,"1. <100mi",IF(${col("min_miles")}{r}<500,"2. 100-500mi",'
   f'IF(${col("min_miles")}{r}<1000,"3. 500-1000mi",IF(${col("min_miles")}{r}<2000,"4. 1000-2000mi","5. 2000+mi")))))'),None,
   'Freight banding. Unknown when neither plant distance was computed - previously these fell into "<100mi".')
dv('is_prevailing_wage','Labor',lambda r:f'=IF(${col("Wage Type")}{r}="Prevailing Wage",1,0)','0','Davis-Bacon style premium')
dv('is_union','Labor',lambda r:f'=IF(${col("Wage Type")}{r}="Union Wage",1,0)','0','Union labor')
dv('no_erection_scope','Labor',lambda r:(f'=IF(OR(${col("Wage Type")}{r}="",${col("Wage Type")}{r}="No Erection Included",'
   f'${col("Wage Type")}{r}="Erection Advisor Only"),1,0)'),'0','KEY: predicts Construction Price = 0')
dv('has_erection_scope','Labor',lambda r:f'=1-${col("no_erection_scope")}{r}','0','Erection included')
dv('labor_rate_index','Labor',lambda r:(f'=IF(${col("is_prevailing_wage")}{r}=1,1.65,IF(${col("is_union")}{r}=1,1.65,'
   f'IF(${col("no_erection_scope")}{r}=1,0,1)))'),'0.00','Observed construction $/sqft premium vs non-union (see Data_Research)')
dv('deck_is_dome','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Dome",${col("Deck Style")}{r})),1,0)','0','Aluminum geodesic dome')
dv('deck_is_open_top','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Open-Top",${col("Deck Style")}{r})),1,0)','0','No roof scope')
dv('deck_has_rafters','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Rafter",${col("Deck Style")}{r})),1,0)','0','Rafter-supported roof')
dv('deck_is_selfsupp','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Self Supported",${col("Deck Style")}{r})),1,0)','0','Self-supported roof')
dv('deck_is_centersupp','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Center Supported",${col("Deck Style")}{r})),1,0)','0','Center column support')
dv('floor_is_embedded_ring','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Embedded Ring",${col("Floor Style")}{r})),1,0)','0','Embedded ring foundation')
dv('floor_is_flat','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Flat",${col("Floor Style")}{r})),1,0)','0','Flat floor')
dv('floor_is_sloped','Options',lambda r:f'=IF(ISNUMBER(SEARCH("Slope",${col("Floor Style")}{r})),1,0)','0','Sloped floor')
dv('floor_is_special','Options',lambda r:f'=IF(${col("Floor Style")}{r}="Special",1,0)','0','Non-standard floor')
dv('is_insulated','Options',lambda r:f'=IF((N(${col("Y Insul Material")}{r})+N(${col("Y Insul Construction")}{r}))>1,1,0)','0',
   'Insulation scope selected - app supplies this at predict time')
dv('tank_count_in_quoterev','Context',lambda r:f'=COUNTIFS(Clean_Data!${CD["Quote-Rev"]}${CF}:${CD["Quote-Rev"]}${CLST},${col("Quote-Rev")}{r})','0',
   'Tanks on the same quote revision')
dv('is_multi_tank_quote','Context',lambda r:f'=IF(${col("tank_count_in_quoterev")}{r}>1,1,0)','0','Multi-tank job')
dv('months_since_2020','Context',lambda r:f'=IF(${col("Due Date")}{r}="","",(${col("Due Date")}{r}-DATE(2020,1,1))/30.44)','#,##0.0',
   'Time index - captures steel-price escalation (~7%/yr observed)')
dv('is_firm_bid','Context',lambda r:f'=IF(${col("Bid Type")}{r}="Firm",1,0)','0','Firm vs budget quote')
# ---------- per-bucket scope gates (stage-1 labels for the two-stage models)
dv('inc_material','Scope Gate',lambda r:f'=IF(N(${col("Y Material")}{r})>1,1,0)','0',
   'Stage-1 label: is the material bucket in scope for this tank')
dv('inc_fabrication','Scope Gate',lambda r:f'=IF(N(${col("Y Fabrication")}{r})>1,1,0)','0',
   'Stage-1 label: is the fabrication bucket in scope')
dv('inc_construction','Scope Gate',lambda r:f'=IF(N(${col("Y Construction")}{r})>1,1,0)','0',
   'Stage-1 label: is erection/construction in scope (34.6% are not)')
dv('inc_insul_material','Scope Gate',lambda r:f'=IF(N(${col("Y Insul Material")}{r})>1,1,0)','0',
   'Stage-1 label: insulation material in scope (~22%)')
dv('inc_insul_construction','Scope Gate',lambda r:f'=IF(N(${col("Y Insul Construction")}{r})>1,1,0)','0',
   'Stage-1 label: insulation construction in scope (~21%)')
dv('scope_class','Scope Gate',lambda r:(f'=IF(AND(${col("inc_material")}{r}=1,${col("inc_construction")}{r}=1),"tank_quote",'
   f'IF(AND(${col("inc_material")}{r}=1,${col("inc_construction")}{r}=0),"materials_only",'
   f'IF(AND(${col("inc_material")}{r}=0,${col("inc_construction")}{r}=1),"construction_only","unknown")))'),None,
   'What the quote actually covers. 35% of tanks are materials-only.')
dv('construction_scope_conflict','Scope Gate',lambda r:(f'=IF(OR(AND(${col("no_erection_scope")}{r}=1,${col("inc_construction")}{r}=1),'
   f'AND(${col("no_erection_scope")}{r}=0,${col("inc_construction")}{r}=0)),1,0)'),'0',
   'Wage type and construction price disagree about erection scope - review these rows')

# ---------- reconciliation / accounting QA
# ---------- money reconciliation.
# Verified identities (4,423/4,423 rows, max relative error 7.1e-07):
#     Proposal Total = bucket sum + Total Tax       (carries NO freight)
#     Total Price    = bucket sum + Freight Price   (carries NO tax)
#     grand total    = bucket sum + Freight + Tax   (exists in NEITHER column)
# Proposal Total and Total Price are two different subtotals. The earlier reading - that
# 39% of rows carried an unexplained +4.7% adder - was sales tax: the correlation between
# (Proposal Total - bucket sum) and Total Tax is 1.0000000000, and every one of those rows
# is a taxed row while every tying row has zero tax.
dv('Ref Grand Total','Reference (not a target)',lambda r:(f'=N(${col("Y Bucket Sum")}{r})+N(${col("Ref Freight")}{r})'
   f'+N(${col("Ref Tax")}{r})'),'$#,##0',
   'buckets + freight + tax. The real all-in number, which no single source column holds.')
dv('recon_proposal_diff','Reconciliation',lambda r:(f'=N(${col("Ref Proposal Total")}{r})-(N(${col("Y Bucket Sum")}{r})'
   f'+N(${col("Ref Tax")}{r}))'),'$#,##0.00',
   'Proposal Total minus (buckets + tax). Zero on every row - this is the identity, not a check of scope.')
dv('recon_proposal_ok','Reconciliation',lambda r:f'=IF(ABS(${col("recon_proposal_diff")}{r})<=1,1,0)','0',
   'Proposal Total ties to buckets + tax within $1')
dv('recon_total_diff','Reconciliation',lambda r:(f'=N(${col("Ref Total Price")}{r})-(N(${col("Y Bucket Sum")}{r})'
   f'+N(${col("Ref Freight")}{r}))'),'$#,##0.00',
   'Total Price minus (buckets + freight). Zero on every row.')
dv('recon_total_ok','Reconciliation',lambda r:f'=IF(ABS(${col("recon_total_diff")}{r})<=1,1,0)','0',
   'Total Price ties to buckets + freight within $1')
dv('is_taxed','Reconciliation',lambda r:f'=IF(N(${col("Ref Tax")}{r})>0,1,0)','0',
   'Sales tax applied. Follows jurisdiction, not estimator: 57.9% of US rows, 0% of Mexican/'
   'Peruvian/Chilean/Argentine rows. Diagnostic only - it is derived from a reference column.')

# ---------- additional geometry / scale
dv('liquid_height_ft','Geometry',lambda r:f'=MAX(0,${col("Height (ft)")}{r}-${col("freeboard_ft")}{r})','#,##0.00',
   'Shell height less freeboard - the wetted height')
dv('liquid_volume_gal','Geometry',lambda r:f'=PI()*(${col("Diameter (ft)")}{r}/2)^2*${col("liquid_height_ft")}{r}*7.48052','#,##0',
   'Volume to the liquid line - closer to usable capacity than gross volume')
dv('wetted_shell_area_sqft','Geometry',lambda r:f'=PI()*${col("Diameter (ft)")}{r}*${col("liquid_height_ft")}{r}','#,##0',
   'Shell area below the liquid line - drives plate thickness')
dv('qty_x_shell_area','Geometry',lambda r:f'=${col("Quantity")}{r}*${col("shell_area_sqft")}{r}','#,##0',
   'Total fabricated area across the line quantity')
dv('roof_geometry','Options',lambda r:(f'=IF(ISNUMBER(SEARCH("Dome",${col("Deck Style")}{r})),"dome",'
   f'IF(ISNUMBER(SEARCH("Open-Top",${col("Deck Style")}{r})),"open_top",'
   f'IF(ISNUMBER(SEARCH(":12",${col("Deck Style")}{r})),"pitched",'
   f'IF(${col("Deck Style")}{r}="","unknown","flat"))))'),None,'Normalised roof class')
dv('floor_geometry','Options',lambda r:(f'=IF(ISNUMBER(SEARCH("Slope",${col("Floor Style")}{r})),"sloped",'
   f'IF(ISNUMBER(SEARCH("Flat",${col("Floor Style")}{r})),"flat",'
   f'IF(OR(ISNUMBER(SEARCH("Ring",${col("Floor Style")}{r})),ISNUMBER(SEARCH("Base Angle",${col("Floor Style")}{r}))),"no_supplied_floor_plate",'
   f'IF(${col("Floor Style")}{r}="","unknown","special"))))'),None,'Normalised floor class')

# ---------- account / context
dv('quote_balanced_weight','Context',lambda r:f'=1/MAX(1,${col("tank_count_in_quoterev")}{r})','0.000',
   'Sample weight so a 29-tank quote does not outvote 29 single-tank quotes')
dv('company_tank_count','Context',lambda r:f'=COUNTIF({mrng("Company Name")},${col("Company Name")}{r})','#,##0',
   'How many tanks we have quoted this end customer')
dv('customer_tank_count','Context',lambda r:f'=COUNTIF({mrng("Customer Name")},${col("Customer Name")}{r})','#,##0',
   'How many tanks quoted through this rep/dealer')
dv('manager_tank_count','Context',lambda r:f'=COUNTIF({mrng("Sales Manager")},${col("Sales Manager")}{r})','#,##0',
   'Sales manager volume - pricing style proxy')

# leakage-flagged research ratios
dv('rate_price_per_gal','Rate (LEAKAGE - research only)',lambda r:f'=IF(N(${col("capacity_gal")}{r})>0,${col("Y Bucket Sum")}{r}/${col("capacity_gal")}{r},"")','0.00',
   'Derived FROM price - never use as a model feature')
dv('rate_price_per_shellsqft','Rate (LEAKAGE - research only)',lambda r:f'=IF(${col("shell_area_sqft")}{r}>0,${col("Y Bucket Sum")}{r}/${col("shell_area_sqft")}{r},"")','0.00',
   'Derived FROM price - never use as a model feature')
dv('shop_cost','Rate (LEAKAGE - research only)',lambda r:f'=N(${col("Y Material")}{r})+N(${col("Y Fabrication")}{r})','$#,##0',
   'Shop scope (material+fabrication). Derived FROM targets - research only.')
dv('rate_material_per_shellsqft','Rate (LEAKAGE - research only)',lambda r:f'=IF(${col("shell_area_sqft")}{r}>0,${col("Y Material")}{r}/${col("shell_area_sqft")}{r},"")','0.00',
   'Derived FROM price - never use as a model feature')

HEADERS=[s['h'] for s in SPEC]
L.update({h:CL(i+1) for i,h in enumerate(HEADERS)})

def build_ml(wb, keys):
    ML_LAST[0]=len(keys)+2
    ws=wb.create_sheet('ML_Tank_Training')
    # group banner row
    ws.append([s['group'] for s in SPEC])
    ws.append(HEADERS)
    for i,k in enumerate(keys):
        r=i+3
        row=[]
        for s in SPEC:
            h=s['h']
            if h=='Tank Key': row.append(k)
            elif s['kind']=='calc' and h=='Src Row':
                row.append(f'=MATCH($A{r},Clean_Data!${CD["Tank Key"]}${CF}:${CD["Tank Key"]}${CLST},0)+1')
            elif s['kind'] in ('pull','target'):
                src=s['f']
                row.append(f'=INDEX(Clean_Data!${CD[src]}:${CD[src]},$B{r})')
            else:
                fn=next(d[2] for d in D if d[0]==h)
                row.append(fn(r))
        ws.append(row)
    # styling
    gfill=PatternFill('solid',fgColor='8EA9DB'); hfill=PatternFill('solid',fgColor='1F3864')
    tfill=PatternFill('solid',fgColor='C6501E'); lfill=PatternFill('solid',fgColor='9C3B3B')
    for i,s in enumerate(SPEC):
        c=i+1
        g=ws.cell(1,c); g.value=s['group']
        g.fill=tfill if s['group']=='TARGET' else (lfill if 'LEAKAGE' in s['group'] else gfill)
        g.font=Font(name=FONT,bold=True,color='FFFFFF',size=8)
        g.alignment=Alignment(horizontal='center',wrap_text=True)
        hc=ws.cell(2,c); hc.fill=hfill; hc.font=Font(name=FONT,bold=True,color='FFFFFF',size=9)
        hc.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        if s['note']: hc.comment=None
        ws.column_dimensions[CL(c)].width=14
        if s['fmt']:
            for r in range(3,len(keys)+3):
                ws.cell(r,c).number_format=s['fmt']
    ws.freeze_panes='C3'
    ws.row_dimensions[1].height=26; ws.row_dimensions[2].height=34
    return ws

if __name__=='__main__':
    src=pd.read_excel('v5_core.xlsx',sheet_name='Clean_Data')
    keys=src.loc[pd.to_numeric(src['ML Training Row'],errors='coerce')==1,'Tank Key'].tolist()
    print('training keys:',len(keys),'unique:',len(set(keys)))
    wb=openpyxl.load_workbook('v5_core.xlsx')
    build_ml(wb,keys)
    wb.save('v5_ml.xlsx')
    print('saved v5_ml.xlsx; ML columns:',len(HEADERS))
    for s in SPEC: print(f"  {L[s['h']]:>3} [{s['group'][:22]:22s}] {s['h']}")
