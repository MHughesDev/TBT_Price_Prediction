"""V5 builder - part 1: Raw_Import (verbatim source) + Clean_Data (fully traceable formulas)."""
import pandas as pd, numpy as np, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from prep import JUNK_WORDS
import geo_crosswalk as GEO

RAW = pd.read_csv('archive.csv', dtype=str).fillna('')
N = len(RAW)
FIRST, LAST = 2, N + 1

RAWCOLS = {h: CL(i + 1) for i, h in enumerate(RAW.columns)}
def R(h, r): return f"Raw_Import!{RAWCOLS[h]}{r}"

FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)

def style_header(ws, ncols, freeze='A2'):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30

def txt(h, r):
    s = R(h, r)
    return (f'=IF(OR(TRIM({s})="",LOWER(TRIM({s}))="null",LOWER(TRIM({s}))="none",'
            f'LOWER(TRIM({s}))="n/a",LOWER(TRIM({s}))="na",TRIM({s})="-"),"",TRIM({s}))')

def numf(h, r):
    s = R(h, r)
    return f'=IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({s},"$",""),",",""),"%","")),"")'

# ---- Clean_Data column headers, in order ----
HEADERS = [
 'QuoteNumber','RevisionNumber','Tank Name','Quote-Rev','Key Dup Index','Tank Key',
 'Bid Type','Status','Row Outcome','Max Revision','Is First Row Of Rev','Revision Count',
 'Is Current Revision','Won Rows On Quote','Lost Rows On Quote','Outcome',
 'Due Date','Due Year','Project Name','Company Name','Customer Name','Sales Manager','Sales Rep',
 'Country','State','City','Wage Type','Miles TBT','Miles GT','Ss','S1',
 'Quantity','Material','Use Type','Deck Style','Floor Style','Diameter (ft)','Height (ft)','Freeboard (in)',
 'Capacity Unit','Capacity Value','Margin (%)','Contingency (%)','Insul Margin (%)','Insul Contingency (%)',
 'Material Price','Fabrication Price','Construction Price','Insul Material Price','Insul Construction Price',
 'Bucket Sum','Total Tax','Freight Price','Proposal Total','Total Price',
 'DQ Has Geometry','DQ Has Capacity','DQ Positive Target','DQ Price Sane','DQ Looks Non-Tank','DQ Unfinished',
 'Geom Volume Gal','DQ Quote Id Numeric','DQ Fill Ratio Valid','DQ Seismic Consistent',
 'DQ Aspect Plausible','DQ Freight Present','DQ Miles GT Zero',
 'ML Eligible','Reporting Eligible',
 'Quote Group ID','Quote Suffix Raw','Quote Suffix','Quote Variant Type','Is Scope Variant',
 'Suffix Revision','Effective Revision','Job Key','Job Max Revision','Is Job Firmest',
 'Job Tank Key','Job Tank First Occurrence','ML Training Row','Reporting Row',
 # --- P0-2 geographic normalization. Keyed on Country|State, not State alone: ten raw
 # --- State values (BC, CA, Santiago, VA, NM, TX, MA, Guam, Ciudad de Mexico, blank)
 # --- resolve to different places depending on the country. See geo_crosswalk.py.
 'Geo Key','Geo Row','Country Normalized','Country Name','State Normalized',
 'Region Code','Geo Match Level','DQ Geo Resolved',
 # --- P0-4 segment-aware outlier bounds. The old DQ Price Sane is a single global
 # --- $1K-$50M band, which calls a $2M price normal for a 1M-gal tank and normal for a
 # --- 5-ft one. These columns rank each row's $/shell-sqft WITHIN its own
 # --- material x scope_class segment and flag the tails. Rate Per Shell SqFt is derived
 # --- from price: it is a data-quality diagnostic ONLY and must never be pulled into
 # --- ML_Tank_Training as a feature (it would leak the target).
 'Shell Area SqFt','Rate Per Shell SqFt','Scope Class','Price Segment',
 'Segment N','Segment Rank Below','Segment Pctile','DQ Price Sane Segment',
 # --- P1-2 revision counting at JOB grain. 'Revision Count' above counts revisions of a
 # --- QuoteNumber; a job (Quote Group ID + scope suffix) can span several quote numbers,
 # --- and 67 quotes have gaps in their revision sequence, so max-min+1 would overcount.
 # --- These two count DISTINCT effective revisions per job, gaps and all.
 'Job Rev First Row','Job Revision Count',
 # --- P1-6 cleaning pass.
 # Canonicalization: 4 company spellings and 2 customer spellings collapse (case, commas,
 # trailing periods). Small in absolute terms, but cardinality is what matters for encoding
 # a 937-level field, so the canonical name is the one to group and encode on.
 'Company Key','Company Canonical','Customer Key','Customer Canonical',
 # Tank Name scope markers. Unlike Scope Class, these are parsed from the NAME the estimator
 # typed and are NOT derived from any price, so they are safe to use as model features.
 # 449 rows are numbered options and ~165 name a roof, deck, floor or demo/replacement -
 # partial-scope work where shell area is not the price driver.
 'Name Has Option','Name Partial Scope']
C = {h: CL(i + 1) for i, h in enumerate(HEADERS)}   # name -> column letter

def formula(h, r):
    c = C
    if h=='QuoteNumber':      return f'=TRIM({R("Quote #",r)})'
    if h=='RevisionNumber':   return f'=IFERROR(VALUE({R("Revision #",r)}),0)'
    if h=='Tank Name':        return f'=IF(TRIM({R("Tank Name",r)})="","Tank ?",TRIM({R("Tank Name",r)}))'
    if h=='Quote-Rev':        return f'=${c["QuoteNumber"]}{r}&"-R"&${c["RevisionNumber"]}{r}'
    if h=='Key Dup Index':    return (f'=COUNTIFS(${c["Quote-Rev"]}$2:${c["Quote-Rev"]}{r},${c["Quote-Rev"]}{r},'
                                      f'${c["Tank Name"]}$2:${c["Tank Name"]}{r},${c["Tank Name"]}{r})')
    if h=='Tank Key':         return (f'=${c["Quote-Rev"]}{r}&"-"&${c["Tank Name"]}{r}&'
                                      f'IF(${c["Key Dup Index"]}{r}>1,"#"&TEXT(${c["Key Dup Index"]}{r},"0"),"")')
    if h=='Bid Type':         return txt('Bid Type', r)
    if h=='Status':           return f'=IF(TRIM({R("Status",r)})="","",PROPER(TRIM({R("Status",r)})))'
    if h=='Row Outcome':      return f'=IF(${c["Status"]}{r}="Won","Won",IF(${c["Status"]}{r}="Lost","Lost","Open"))'
    if h=='Max Revision':     return (f'=_xlfn.MAXIFS(${c["RevisionNumber"]}${FIRST}:${c["RevisionNumber"]}${LAST},'
                                      f'${c["QuoteNumber"]}${FIRST}:${c["QuoteNumber"]}${LAST},${c["QuoteNumber"]}{r})')
    if h=='Is First Row Of Rev': return f'=IF(COUNTIFS(${c["Quote-Rev"]}$2:${c["Quote-Rev"]}{r},${c["Quote-Rev"]}{r})=1,1,0)'
    if h=='Revision Count':   return (f'=SUMIFS(${c["Is First Row Of Rev"]}${FIRST}:${c["Is First Row Of Rev"]}${LAST},'
                                      f'${c["QuoteNumber"]}${FIRST}:${c["QuoteNumber"]}${LAST},${c["QuoteNumber"]}{r})')
    if h=='Is Current Revision': return f'=IF(${c["RevisionNumber"]}{r}=${c["Max Revision"]}{r},TRUE,FALSE)'
    if h=='Won Rows On Quote':return (f'=COUNTIFS(${c["QuoteNumber"]}${FIRST}:${c["QuoteNumber"]}${LAST},${c["QuoteNumber"]}{r},'
                                      f'${c["Status"]}${FIRST}:${c["Status"]}${LAST},"Won")')
    if h=='Lost Rows On Quote':return (f'=COUNTIFS(${c["QuoteNumber"]}${FIRST}:${c["QuoteNumber"]}${LAST},${c["QuoteNumber"]}{r},'
                                      f'${c["Status"]}${FIRST}:${c["Status"]}${LAST},"Lost")')
    if h=='Outcome':          return (f'=IF(${c["Won Rows On Quote"]}{r}>0,"Won",'
                                      f'IF(${c["Lost Rows On Quote"]}{r}>0,"Lost","Open"))')
    if h=='Due Date':         return f'=IFERROR(DATEVALUE({R("Due Date",r)}),"")'
    if h=='Due Year':         return f'=IF(${c["Due Date"]}{r}="","",YEAR(${c["Due Date"]}{r}))'
    if h=='Project Name':     return txt('Project Name', r)
    if h=='Company Name':     return txt('Company Name', r)
    if h=='Customer Name':
        s = R('Customer Name', r)
        return (f'=IF(OR(TRIM({s})="",LOWER(TRIM({s}))="null",LOWER(TRIM({s}))="none",LOWER(TRIM({s}))="n/a",'
                f'LOWER(TRIM({s}))="na",TRIM({s})="-"),"(No Customer Listed)",TRIM({s}))')
    if h=='Sales Manager':    return txt('Sales Manager', r)
    if h=='Sales Rep':        return txt('Sales Rep', r)
    if h=='Country':          return txt('Country', r)
    if h=='State':            return txt('State', r)
    if h=='City':             return f'=IF(TRIM({R("City",r)})="","",TRIM(SUBSTITUTE(TRIM({R("City",r)}),",","")))'
    if h=='Wage Type':        return txt('Wage Type', r)
    if h=='Miles TBT':        return numf('Miles to Site (From TBT)', r)
    if h=='Miles GT':         return numf('Miles to Site (From GT)', r)
    if h=='Ss':               return numf('Ss', r)
    if h=='S1':               return numf('S1', r)
    if h=='Quantity':         return f'=IFERROR(VALUE({R("Quantity",r)}),1)'
    if h=='Material':         return txt('Material', r)
    if h=='Use Type':         return txt('Use Type', r)
    if h=='Deck Style':       return txt('Deck Style', r)
    if h=='Floor Style':      return txt('Floor Style', r)
    if h=='Diameter (ft)':    return numf('Diameter (ft)', r)
    if h=='Height (ft)':      return numf('Height (ft)', r)
    if h=='Freeboard (in)':   return numf('Freeboard (in)', r)
    if h=='Capacity Unit':
        s = R('Usable Capacity', r)
        return f'=IF(ISNUMBER(SEARCH("ton",{s})),"tons",IF(ISNUMBER(SEARCH("gal",{s})),"gal",""))'
    if h=='Capacity Value':
        s = R('Usable Capacity', r)
        return (f'=IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(LOWER({s}),",",""),'
                f'"gal",""),"tons",""),"ton","")),"")')
    if h=='Margin (%)':       return numf('Margin (%)', r)
    if h=='Contingency (%)':  return numf('Contingency (%)', r)
    if h=='Insul Margin (%)': return numf('Insulation Margin (%)', r)
    if h=='Insul Contingency (%)': return numf('Insulation Contingency (%)', r)
    if h=='Material Price':   return numf('Material Price', r)
    if h=='Fabrication Price':return numf('Fabrication Price', r)
    if h=='Construction Price':return numf('Construction Price', r)
    if h=='Insul Material Price': return numf('Insulation Material Price', r)
    if h=='Insul Construction Price': return numf('Insulation Construction Price', r)
    if h=='Bucket Sum':       return f'=SUM(${c["Material Price"]}{r}:${c["Insul Construction Price"]}{r})'
    if h=='Total Tax':        return numf('Total Tax', r)
    if h=='Freight Price':    return numf('Freight Price', r)
    if h=='Proposal Total':   return numf('Proposal Total', r)
    if h=='Total Price':      return numf('Total Price', r)
    if h=='DQ Has Geometry':  return f'=IF(AND(N(${c["Diameter (ft)"]}{r})>0,N(${c["Height (ft)"]}{r})>0),1,0)'
    # P1-6: 44 source rows carry a negative usable capacity ("-937.376tons"). A negative
    # volume cannot be corrected into a real one, so it is treated as MISSING everywhere
    # downstream; the signed value stays in Capacity Value so the error remains traceable.
    if h=='DQ Has Capacity':  return f'=IF(N(${c["Capacity Value"]}{r})>0,1,0)'
    if h=='DQ Positive Target':return f'=IF(${c["Bucket Sum"]}{r}>0,1,0)'
    if h=='DQ Price Sane':    return (f'=IF(AND(N(${c["Proposal Total"]}{r})>=1000,'
                                      f'N(${c["Proposal Total"]}{r})<=50000000),1,0)')
    if h=='DQ Looks Non-Tank':return (f'=IF(SUMPRODUCT(--ISNUMBER(SEARCH(Ref_JunkWords,LOWER(${c["Tank Name"]}{r}&" "&'
                                      f'${c["QuoteNumber"]}{r}&" "&${c["Use Type"]}{r}))))>0,1,0)')
    if h=='DQ Unfinished':    return f'=IF(OR(${c["Status"]}{r}="Unfinished",${c["Status"]}{r}="Bid Review"),1,0)'
    if h=='Geom Volume Gal':
        return (f'=IF(AND(N(${c["Diameter (ft)"]}{r})>0,N(${c["Height (ft)"]}{r})>0),'
                f'PI()*(${c["Diameter (ft)"]}{r}/2)^2*${c["Height (ft)"]}{r}*7.48052,0)')
    if h=='DQ Quote Id Numeric':
        return f'=IFERROR(IF(VALUE(LEFT($A{r},7))>0,1,0),0)'
    if h=='DQ Fill Ratio Valid':
        cap=f'${c["Capacity Value"]}{r}'; gv=f'${c["Geom Volume Gal"]}{r}'
        return (f'=IF(${c["Capacity Unit"]}{r}<>"gal",1,IF({gv}<=0,1,'
                f'IF(N({cap})/{gv}<=1.05,1,0)))')
    if h=='DQ Seismic Consistent':
        return f'=IF(N(${c["S1"]}{r})<=N(${c["Ss"]}{r}),1,0)'
    if h=='DQ Aspect Plausible':
        return (f'=IF(N(${c["Diameter (ft)"]}{r})<=0,0,'
                f'IF(AND(N(${c["Height (ft)"]}{r})/N(${c["Diameter (ft)"]}{r})>=0.05,'
                f'N(${c["Height (ft)"]}{r})/N(${c["Diameter (ft)"]}{r})<=6),1,0))')
    if h=='DQ Freight Present':
        return (f'=IF(OR(N(${c["Freight Price"]}{r})>0,'
                f'MIN(N(${c["Miles TBT"]}{r}),N(${c["Miles GT"]}{r}))<=100),1,0)')
    if h=='DQ Miles GT Zero':
        return f'=IF(N(${c["Miles GT"]}{r})=0,1,0)'
    # P0-4: the global band stays as a backstop, the segment rule is added on top.
    if h=='ML Eligible':      return (f'=IF(AND(${c["DQ Has Geometry"]}{r}=1,${c["DQ Positive Target"]}{r}=1,'
                                      f'${c["DQ Price Sane"]}{r}=1,${c["DQ Price Sane Segment"]}{r}=1,'
                                      f'${c["DQ Looks Non-Tank"]}{r}=0,'
                                      f'${c["DQ Unfinished"]}{r}=0),1,0)')
    if h=='Quote Group ID':
        return f'=IFERROR(IF(VALUE(LEFT($A{r},7))>0,LEFT($A{r},7),$A{r}),$A{r})'
    if h=='Quote Suffix Raw':
        return f'=IF(${c["Quote Group ID"]}{r}=$A{r},"",TRIM(MID($A{r},8,99)))'
    if h=='Quote Suffix':
        q=f'${c["Quote Suffix Raw"]}{r}'
        return f'=IF(LEFT({q},1)="-",TRIM(MID({q},2,99)),{q})'
    if h=='Quote Variant Type':
        q=f'${c["Quote Suffix"]}{r}'
        t=('=IF(@="","bare",'
           'IF(ISNUMBER(SEARCH("as sold",@)),"status_snapshot",'
           'IF(ISNUMBER(SEARCH("as approved",@)),"status_snapshot",'
           'IF(ISNUMBER(SEARCH("co#",@)),"change_order",'
           'IF(AND(LEFT(@,2)="CO",ISNUMBER(IFERROR(VALUE(MID(@,3,2)),"x"))),"change_order",'
           'IF(AND(LEFT(@,2)="CO",ISNUMBER(IFERROR(VALUE(MID(@,3,1)),"x"))),"change_order",'
           'IF(ISNUMBER(SEARCH("option",@)),"option",'
           'IF(ISNUMBER(SEARCH("only",@)),"scope_variant",'
           'IF(ISNUMBER(SEARCH("materials",@)),"scope_variant",'
           'IF(ISNUMBER(SEARCH("eng",@)),"scope_variant",'
           'IF(AND(LEFT(@,1)="R",ISNUMBER(IFERROR(VALUE(MID(@,2,1)),"x"))),"revision",'
           '"other")))))))))))')
        return t.replace('@',q)
    if h=='Is Scope Variant':
        v=f'${c["Quote Variant Type"]}{r}'
        return f'=IF(OR({v}="change_order",{v}="option",{v}="scope_variant",{v}="other"),1,0)'
    if h=='Suffix Revision':
        q=f'${c["Quote Suffix"]}{r}'
        return (f'=IF(${c["Quote Variant Type"]}{r}<>"revision",0,'
                f'IF(AND(MID({q},3,1)<>"",ISNUMBER(SEARCH(MID({q},3,1),"0123456789"))),'
                f'IFERROR(VALUE(MID({q},2,2)),0),IFERROR(VALUE(MID({q},2,1)),0)))')
    if h=='Effective Revision':
        return f'=MAX(${c["RevisionNumber"]}{r},${c["Suffix Revision"]}{r})'
    if h=='Job Key':
        return (f'=${c["Quote Group ID"]}{r}&"|"&IF(${c["Is Scope Variant"]}{r}=1,'
                f'${c["Quote Suffix"]}{r},"")')
    if h=='Job Max Revision':
        return (f'=_xlfn.MAXIFS(${c["Effective Revision"]}${FIRST}:${c["Effective Revision"]}${LAST},'
                f'${c["Job Key"]}${FIRST}:${c["Job Key"]}${LAST},${c["Job Key"]}{r})')
    if h=='Is Job Firmest':
        return f'=IF(${c["Effective Revision"]}{r}=${c["Job Max Revision"]}{r},1,0)'
    if h=='Job Tank Key':
        return f'=${c["Job Key"]}{r}&"|"&${c["Tank Name"]}{r}'
    if h=='Job Tank First Occurrence':
        jt=c["Job Tank Key"]; jf=c["Is Job Firmest"]
        return (f'=IF(AND(${jf}{r}=1,COUNTIFS(${jt}$2:${jt}{r},${jt}{r},'
                f'${jf}$2:${jf}{r},1)=1),1,0)')
    if h=='Reporting Eligible':
        return (f'=IF(AND(${c["DQ Price Sane"]}{r}=1,${c["DQ Looks Non-Tank"]}{r}=0,'
                f'${c["DQ Unfinished"]}{r}=0,${c["DQ Quote Id Numeric"]}{r}=1),1,0)')
    if h=='ML Training Row':
        return (f'=IF(AND(${c["ML Eligible"]}{r}=1,${c["Is Job Firmest"]}{r}=1,'
                f'${c["Job Tank First Occurrence"]}{r}=1),1,0)')
    if h=='Reporting Row':
        return (f'=IF(AND(${c["Reporting Eligible"]}{r}=1,${c["Is Job Firmest"]}{r}=1),1,0)')
    # ---- P0-2 geography. One MATCH per row feeds five INDEXes, the same pattern
    # ---- ML_Tank_Training uses for Src Row. INDEX/MATCH, never XLOOKUP (gotcha 4).
    if h=='Geo Key':
        return f'=UPPER(TRIM(${c["Country"]}{r})&"|"&TRIM(${c["State"]}{r}))'
    if h=='Geo Row':
        return f'=IFERROR(MATCH(${c["Geo Key"]}{r},Ref_GeoKey,0),0)'
    if h in ('Country Normalized','Country Name','State Normalized','Region Code','Geo Match Level'):
        n = GEO.HEADERS.index(h) + 1              # 1-based column inside the crosswalk block
        # T() coerces an empty crosswalk cell to "" instead of the numeric 0 that a bare
        # INDEX returns. Every crosswalk column is text, so T() is lossless here.
        return (f'=IF(${c["Geo Row"]}{r}=0,"",T(INDEX(Ref_GeoTable,${c["Geo Row"]}{r},{n})))')
    if h=='DQ Geo Resolved':
        return f'=IF(${c["Geo Match Level"]}{r}="region",1,0)'
    # ---- P0-4 segment-aware price sanity -----------------------------------
    if h=='Shell Area SqFt':
        return (f'=IF(AND(N(${c["Diameter (ft)"]}{r})>0,N(${c["Height (ft)"]}{r})>0),'
                f'PI()*${c["Diameter (ft)"]}{r}*${c["Height (ft)"]}{r},"")')
    if h=='Rate Per Shell SqFt':
        return (f'=IF(AND(N(${c["Shell Area SqFt"]}{r})>0,N(${c["Bucket Sum"]}{r})>0),'
                f'${c["Bucket Sum"]}{r}/${c["Shell Area SqFt"]}{r},"")')
    if h=='Scope Class':
        m=f'OR(N(${c["Material Price"]}{r})>1,N(${c["Fabrication Price"]}{r})>1)'
        cc=f'N(${c["Construction Price"]}{r})>1'
        return (f'=IF(AND({m},{cc}),"tank_quote",IF(AND({m},NOT({cc})),"materials_only",'
                f'IF(AND(NOT({m}),{cc}),"construction_only","unknown")))')
    if h=='Price Segment':
        return f'=${c["Material"]}{r}&"|"&${c["Scope Class"]}{r}'
    if h=='Segment N':
        # rows in this segment that have a computable rate
        return (f'=COUNTIFS(${c["Price Segment"]}${FIRST}:${c["Price Segment"]}${LAST},${c["Price Segment"]}{r},'
                f'${c["Rate Per Shell SqFt"]}${FIRST}:${c["Rate Per Shell SqFt"]}${LAST},">0")')
    if h=='Segment Rank Below':
        return (f'=IF(${c["Rate Per Shell SqFt"]}{r}="","",'
                f'COUNTIFS(${c["Price Segment"]}${FIRST}:${c["Price Segment"]}${LAST},${c["Price Segment"]}{r},'
                f'${c["Rate Per Shell SqFt"]}${FIRST}:${c["Rate Per Shell SqFt"]}${LAST},'
                f'"<"&${c["Rate Per Shell SqFt"]}{r}))')
    if h=='Segment Pctile':
        return (f'=IF(OR(${c["Segment Rank Below"]}{r}="",${c["Segment N"]}{r}=0),"",'
                f'${c["Segment Rank Below"]}{r}/${c["Segment N"]}{r})')
    # ---- P1-6 canonicalization. Strip case, spaces and the punctuation that actually
    # ---- varies in this data (. , - &). The canonical label is the first spelling seen.
    if h in ('Company Key','Customer Key'):
        src = c['Company Name'] if h=='Company Key' else c['Customer Name']
        e = f'${src}{r}'
        for ch in ('.', ',', '-', '&', ' '):
            e = f'SUBSTITUTE({e},"{ch}","")'
        return f'=UPPER({e})'
    if h in ('Company Canonical','Customer Canonical'):
        key = c['Company Key'] if h=='Company Canonical' else c['Customer Key']
        nm  = c['Company Name'] if h=='Company Canonical' else c['Customer Name']
        return (f'=IF(${key}{r}="","",IFERROR(INDEX(${nm}${FIRST}:${nm}${LAST},'
                f'MATCH(${key}{r},${key}${FIRST}:${key}${LAST},0)),${nm}{r}))')
    if h=='Name Has Option':
        n=f'LOWER(${c["Tank Name"]}{r})'
        return f'=IF(ISNUMBER(SEARCH("option",{n})),1,0)'
    if h=='Name Partial Scope':
        n=f'LOWER(${c["Tank Name"]}{r})'
        terms=['roof','deck','floor','demo','replacement','rings','shell course','nozzle','ladder']
        tests='+'.join(f'--ISNUMBER(SEARCH("{t}",{n}))' for t in terms)
        return f'=IF(({tests})>0,1,0)'
    if h=='Job Rev First Row':
        jk=c["Job Key"]; er=c["Effective Revision"]
        return (f'=IF(COUNTIFS(${jk}${FIRST}:${jk}{r},${jk}{r},'
                f'${er}${FIRST}:${er}{r},${er}{r})=1,1,0)')
    if h=='Job Revision Count':
        return (f'=SUMIFS(${c["Job Rev First Row"]}${FIRST}:${c["Job Rev First Row"]}${LAST},'
                f'${c["Job Key"]}${FIRST}:${c["Job Key"]}${LAST},${c["Job Key"]}{r})')
    if h=='DQ Price Sane Segment':
        # Pass when no rate is computable (geometry is judged by DQ Has Geometry) or when
        # the segment is too small to place a tail. Otherwise flag the 1.5% / 99% tails.
        return (f'=IF(${c["Segment Pctile"]}{r}="",1,IF(${c["Segment N"]}{r}<30,1,'
                f'IF(OR(${c["Segment Pctile"]}{r}<0.015,${c["Segment Pctile"]}{r}>0.99),0,1)))')
    raise KeyError(h)

def build(wb):
    ws = wb.create_sheet('Raw_Import')
    ws.append(list(RAW.columns))
    for row in RAW.itertuples(index=False):
        ws.append(list(row))
    style_header(ws, len(RAW.columns))
    for c in range(1, len(RAW.columns) + 1):
        ws.column_dimensions[CL(c)].width = 16

    ws = wb.create_sheet('Clean_Data')
    ws.append(HEADERS)
    for r in range(FIRST, LAST + 1):
        ws.append([formula(h, r) for h in HEADERS])
    style_header(ws, len(HEADERS))
    for c in range(1, len(HEADERS) + 1):
        ws.column_dimensions[CL(c)].width = 15
    # Due Date is a DATEVALUE serial. Without an explicit date format it renders as
    # "45531" and pandas reads it as int64. The delivered V5 only looked right because
    # LibreOffice silently inferred a date format during recalc; Excel does not, so the
    # format has to be stated here. build_dash.py already does this for Fact_Quote.
    dd = C['Due Date']
    for r in range(FIRST, LAST + 1):
        ws[f'{dd}{r}'].number_format = 'yyyy-mm-dd'

    ref = wb.create_sheet('Ref_Lists')
    ref['A1'] = 'Junk / non-tank keywords (drives Clean_Data "DQ Looks Non-Tank")'
    ref['A1'].font = Font(name=FONT, bold=True)
    words = JUNK_WORDS   # canonical list lives in prep.py so Excel and the pandas reference agree
    for i, w in enumerate(words):
        ref.cell(2 + i, 1, w)
    wb.defined_names.add(DefinedName('Ref_JunkWords', attr_text=f"Ref_Lists!$A$2:$A${1+len(words)}"))

    # ---- P0-2 geographic crosswalk (Country|State -> ISO region + corrected country).
    # Columns C..K; column C is the lookup key, C..K the table INDEX reads from.
    grows = GEO.rows()
    ref['C1'] = 'Geographic crosswalk - raw Country|State -> normalized region + country'
    ref['C1'].font = Font(name=FONT, bold=True)
    for j, hh in enumerate(GEO.HEADERS):
        cell = ref.cell(2, 3 + j, hh)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
    for i, row in enumerate(grows):
        for j, v in enumerate(row):
            ref.cell(3 + i, 3 + j, v)
    first, last = 3, 2 + len(grows)
    ref.column_dimensions['C'].width = 26
    for j in range(1, len(GEO.HEADERS)):
        ref.column_dimensions[CL(3 + j)].width = 20
    ref.column_dimensions[CL(3 + len(GEO.HEADERS) - 1)].width = 70   # Note
    wb.defined_names.add(DefinedName('Ref_GeoKey',
        attr_text=f"Ref_Lists!$C${first}:$C${last}"))
    wb.defined_names.add(DefinedName('Ref_GeoTable',
        attr_text=f"Ref_Lists!$C${first}:${CL(2+len(GEO.HEADERS))}${last}"))
    ref.freeze_panes = 'D3'
    return C

if __name__ == '__main__':
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    build(wb)
    wb.save('v5_core.xlsx')
    print('saved. Clean_Data columns:')
    for h in HEADERS: print(f'  {C[h]:>3} {h}')
